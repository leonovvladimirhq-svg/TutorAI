"""Отчёты для наставника — .docx из реальных данных.

Два вида:
- карта целей одного студента (build_student_report);
- общий отчёт по всем студентам наставника (build_group_report): сводная таблица
  по группе + карта каждого студента отдельным разделом. Собирается одной кнопкой,
  когда студенты заполнили цели и прошли подведение итогов.

Структура карты по согласованному макету: шапка, сводка целей, разбор по каждой цели
(SMART + решение наставника + рефлексия), закономерности (сначала слова студента,
потом сводка ИИ), критерии оценивания для наставника. Степень достижения целей
не оценивается — предмет оценки только качество целей и глубина рефлексии.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import crud
from app.services.reflection import OUTCOME_LABELS, QUESTION_LABELS

ACCENT = RGBColor(0x5B, 0x4B, 0x9E)
GREY = RGBColor(0x77, 0x77, 0x77)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OK = RGBColor(0x1E, 0x84, 0x49)
MID = RGBColor(0xB7, 0x79, 0x1F)
BAD = RGBColor(0xC0, 0x39, 0x2B)

_STATUS_RU = {"draft": "черновик", "active": "активна", "done": "достигнута", "dropped": "снята"}
_CONFIRM_RU = {"pending": ("ждёт подтверждения", MID), "confirmed": ("подтверждена", OK), "rejected": ("возвращена на доработку", BAD)}
_OUTCOME_COLOR = {"achieved": OK, "partial": MID, "not_achieved": BAD, "irrelevant": GREY}

CRITERIA = [
    "Цели полноценные, ясные по SMART, чёткие и понятные.",
    "Есть ясный план выполнения — что конкретно студент будет делать.",
    "В отчёте есть отметки о движении по всем целям и действиям.",
    "Рефлексия глубокая и полноценная: видно, что студент задумывался о происходящем.",
    "Степень достижения цели НЕ оценивается: студент может не достичь ничего, но качественно отрефлексировать результат.",
]


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _run(p, text: str, bold=False, italic=False, color: RGBColor | None = None, size: int | None = None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if size is not None:
        r.font.size = Pt(size)
    return r


def _para(doc, text: str = "", **kw):
    p = doc.add_paragraph()
    if text:
        _run(p, text, **kw)
    p.paragraph_format.space_after = Pt(6)
    return p


def _kv_table(doc, rows: list[tuple[str, str]], key_fill="F4F2FA", key_w=5.0, val_w=12.0):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Cm(key_w), Cm(val_w)
        _shade(cells[0], key_fill)
        _run(cells[0].paragraphs[0], k, bold=True, color=GREY, size=9.5)
        _run(cells[1].paragraphs[0], v or "—", size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def _header_row(t, headers: list[str]) -> None:
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        _shade(c, "5B4B9E")
        _run(c.paragraphs[0], h, bold=True, color=WHITE, size=9.5)


def _fmt_dt(dt: datetime | None) -> str:
    return dt.strftime("%d.%m.%Y") if dt else "—"


def _new_document() -> Document:
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    return doc


def _safe_name(name: str) -> str:
    return "".join(ch for ch in name if ch.isalnum() or ch in " _-").strip().replace(" ", "_") or "student"


def _footer(doc) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    _run(p, "Сформировано автоматически системой TutorAI. Данные обрабатываются в соответствии "
            "с согласием студента (152-ФЗ).", italic=True, color=GREY, size=8.5)


async def _load_student(session: AsyncSession, app_user) -> dict:
    """Всё, что нужно для карты студента, одним словарём."""
    student = await crud.get_student_by_tg(session, app_user.telegram_id)
    goals = await crud.list_goals(session, student.id) if student else []
    refl = await crud.latest_completed_reflection(session, student.id) if student else None
    grefl = {}
    if refl:
        for gr in await crud.list_goal_reflections(session, refl.id):
            grefl[gr.goal_id] = gr
    return {"app_user": app_user, "student": student, "goals": goals, "refl": refl, "grefl": grefl,
            "name": app_user.full_name or f"ID {app_user.telegram_id}"}


async def _render_student(doc, session: AsyncSession, data: dict, *, mentor, goals_deadline, reflection_deadline,
                          title_level: int = 1, prefix: str = "", intro: bool = False) -> None:
    """Карта целей одного студента. В общем отчёте prefix — номер раздела («2.3 »)."""
    app_user, goals, refl, grefl, name = data["app_user"], data["goals"], data["refl"], data["grefl"], data["name"]
    h = 1 if title_level == 1 else title_level

    _kv_table(doc, [
        ("Студент", name),
        ("ID в MAX", str(app_user.telegram_id)),
        ("Наставник", (mentor.full_name or str(mentor.telegram_id)) if mentor else "не назначен"),
        ("Срок подачи целей", goals_deadline or "не задан"),
        ("Срок рефлексии", reflection_deadline or "не задан"),
        ("Целей поставлено", str(len(goals))),
        ("Рефлексия пройдена", _fmt_dt(refl.completed_at) if refl else "ещё нет"),
        ("Отчёт сформирован", datetime.now().strftime("%d.%m.%Y %H:%M")),
    ])
    if intro:
        _intro(doc)

    # --- 1. сводка ---
    doc.add_heading(f"{prefix}1. Цели периода: сводка", level=h)
    if not goals:
        _para(doc, "Студент пока не поставил ни одной цели.", italic=True, color=GREY)
    else:
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        _header_row(t, ["№", "Цель", "Решение наставника", "Итог рефлексии"])
        for n, g in enumerate(goals, 1):
            cells = t.add_row().cells
            _run(cells[0].paragraphs[0], str(n), size=9.5)
            _run(cells[1].paragraphs[0], g.title, size=9.5)
            label, color = _CONFIRM_RU.get(g.confirm_status, (g.confirm_status, GREY))
            _run(cells[2].paragraphs[0], label, bold=True, color=color, size=9.5)
            gr = grefl.get(g.id)
            if g.status == "dropped":
                _run(cells[3].paragraphs[0], OUTCOME_LABELS["irrelevant"], bold=True, color=GREY, size=9.5)
            elif gr:
                _run(cells[3].paragraphs[0], OUTCOME_LABELS.get(gr.outcome, gr.outcome), bold=True,
                     color=_OUTCOME_COLOR.get(gr.outcome, GREY), size=9.5)
            else:
                _run(cells[3].paragraphs[0], "рефлексии ещё нет", color=GREY, size=9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- 2. разбор по целям ---
    doc.add_heading(f"{prefix}2. Разбор по каждой цели", level=h)
    if not goals:
        _para(doc, "Разбирать пока нечего.", italic=True, color=GREY)
    for n, g in enumerate(goals, 1):
        p = doc.add_paragraph()
        _run(p, f"Цель {n}. {g.title}", bold=True, color=ACCENT, size=12)
        p.paragraph_format.space_before = Pt(10)
        _kv_table(doc, [
            ("S — конкретность", g.specific or "—"),
            ("M — измеримость", g.measurable or "—"),
            ("A — достижимость", g.achievable or "—"),
            ("R — значимость", g.relevant or "—"),
            ("T — срок", g.time_bound or "—"),
        ], key_fill="F7F6FC")
        label, color = _CONFIRM_RU.get(g.confirm_status, (g.confirm_status, GREY))
        p = doc.add_paragraph()
        _run(p, "Решение наставника: ", bold=True)
        _run(p, label, bold=True, color=color)
        if g.mentor_comment:
            _run(p, f" — «{g.mentor_comment}»", italic=True)
        if g.status == "dropped":
            p = doc.add_paragraph()
            _run(p, "Актуальность: ", bold=True)
            _run(p, "студент отметил, что цель потеряла актуальность", color=GREY)
            if g.relevance_note:
                _run(p, f" — «{g.relevance_note}»", italic=True)
        gr = grefl.get(g.id)
        if gr:
            p = doc.add_paragraph()
            _run(p, "Итог по словам студента: ", bold=True)
            _run(p, OUTCOME_LABELS.get(gr.outcome, gr.outcome), bold=True, color=_OUTCOME_COLOR.get(gr.outcome, GREY))
            answered = 0
            for key, qlabel in QUESTION_LABELS.items():
                ans = (gr.answers or {}).get(key)
                if ans:
                    answered += 1
                    p = doc.add_paragraph()
                    _run(p, f"{qlabel}: ", bold=True, color=ACCENT, size=10)
                    _run(p, f"«{ans}»", italic=True, size=10)
            if gr.outcome != "irrelevant" and answered < len(QUESTION_LABELS):
                _para(doc, f"Студент ответил на {answered} из {len(QUESTION_LABELS)} вопросов "
                           "(часть пропущена или подведение итогов завершено досрочно).",
                      italic=True, color=GREY, size=9)
        elif g.status != "dropped":
            _para(doc, "Рефлексия по этой цели ещё не проводилась.", italic=True, color=GREY, size=9.5)

    # --- 3. закономерности ---
    doc.add_heading(f"{prefix}3. Закономерности: взгляд студента и сводка системы", level=h)
    if refl and (refl.student_patterns or refl.ai_summary):
        _para(doc, "Что студент увидел сам", bold=True, color=ACCENT, size=11)
        _para(doc, f"«{refl.student_patterns}»" if refl.student_patterns else "Студент не сформулировал закономерности.",
              italic=bool(refl.student_patterns), color=None if refl.student_patterns else GREY)
        _para(doc, "Что добавила система", bold=True, color=ACCENT, size=11)
        _para(doc, refl.ai_summary or "Сводка не сформирована.", color=None if refl.ai_summary else GREY)
    else:
        _para(doc, "Раздел заполнится после того, как студент пройдёт подведение итогов по целям.",
              italic=True, color=GREY)

    # --- 4. критерии ---
    doc.add_heading(f"{prefix}4. Критерии оценивания (для наставника)", level=h)
    _para(doc, "Те же критерии, что применяются к письменному заданию. Отметьте, что видите в материале выше:")
    for c in CRITERIA:
        p = doc.add_paragraph(style="List Bullet")
        _run(p, "☐  " + c, size=10)
    _para(doc, "К разговору с наставником: обсуждать не «почему не сделал», а что именно происходило "
               "на каждом шаге и что студент об этом думает.", italic=True, color=GREY, size=9.5)


def _intro(doc) -> None:
    _para(doc, "Что это за документ", bold=True, color=ACCENT, size=12)
    _para(doc, "Это не ведомость успеваемости: здесь не выставляются оценки и не фиксируются "
               "нарушения. Документ отражает то, что студент поставил себе сам, что из этого "
               "получилось, как он это осмыслил — и какие закономерности за этим видны. "
               "Он передаётся наставнику как основа для разговора, а не как отчёт о проделанной работе.")
    _para(doc, "Степень достижения целей не оценивается. Предмет оценки — качество целей и глубина рефлексии.",
          italic=True, color=GREY, size=9.5)


async def build_student_report(session: AsyncSession, app_user) -> tuple[bytes, str]:
    """Собрать .docx по студенту (запись реестра ролей). Возвращает (bytes, имя файла)."""
    data = await _load_student(session, app_user)
    mentor = await crud.get_app_user_by_tg(session, app_user.mentor_tg) if app_user.mentor_tg else None
    goals_deadline = await crud.get_setting(session, "goals_deadline")
    reflection_deadline = await crud.get_setting(session, "reflection_deadline")

    doc = _new_document()
    p = doc.add_paragraph()
    _run(p, "Индивидуальная карта целей студента", bold=True, color=ACCENT, size=18)
    p = doc.add_paragraph()
    _run(p, "TutorAI · ИИ-наставник · Школа коммуникаций НИУ ВШЭ", color=GREY, size=10)
    p.paragraph_format.space_after = Pt(10)
    await _render_student(doc, session, data, mentor=mentor, goals_deadline=goals_deadline,
                          reflection_deadline=reflection_deadline, intro=True)
    _footer(doc)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue(), f"TutorAI_отчёт_{_safe_name(data['name'])}.docx"


async def build_group_report(session: AsyncSession, mentor_user) -> tuple[bytes, str, int]:
    """Общий .docx по всем студентам наставника. Возвращает (bytes, имя файла, число студентов)."""
    users = await crud.list_students_of_mentor(session, mentor_user.telegram_id)
    goals_deadline = await crud.get_setting(session, "goals_deadline")
    reflection_deadline = await crud.get_setting(session, "reflection_deadline")
    datas = [await _load_student(session, u) for u in users]
    mentor_name = mentor_user.full_name or f"ID {mentor_user.telegram_id}"

    doc = _new_document()
    p = doc.add_paragraph()
    _run(p, "Карты целей студентов: общий отчёт наставника", bold=True, color=ACCENT, size=18)
    p = doc.add_paragraph()
    _run(p, "TutorAI · ИИ-наставник · Школа коммуникаций НИУ ВШЭ", color=GREY, size=10)
    p.paragraph_format.space_after = Pt(10)

    total_goals = sum(len(d["goals"]) for d in datas)
    _kv_table(doc, [
        ("Наставник", mentor_name),
        ("Студентов закреплено", str(len(datas))),
        ("Целей всего", str(total_goals)),
        ("Подтверждено наставником", str(sum(1 for d in datas for g in d["goals"] if g.confirm_status == "confirmed"))),
        ("Прошли подведение итогов", f"{sum(1 for d in datas if d['refl'])} из {len(datas)}"),
        ("Срок подачи целей", goals_deadline or "не задан"),
        ("Срок рефлексии", reflection_deadline or "не задан"),
        ("Отчёт сформирован", datetime.now().strftime("%d.%m.%Y %H:%M")),
    ])
    _intro(doc)

    # --- сводная таблица по группе ---
    doc.add_heading("1. Группа: кто на каком этапе", level=1)
    if not datas:
        _para(doc, "За наставником пока не закреплены студенты.", italic=True, color=GREY)
    else:
        t = doc.add_table(rows=1, cols=6)
        t.style = "Table Grid"
        _header_row(t, ["№", "Студент", "Целей", "Подтв. / возвр. / ждут", "Рефлексия", "Итоги по целям"])
        for n, d in enumerate(datas, 1):
            goals = d["goals"]
            cells = t.add_row().cells
            _run(cells[0].paragraphs[0], str(n), size=9.5)
            _run(cells[1].paragraphs[0], d["name"], bold=True, size=9.5)
            _run(cells[2].paragraphs[0], str(len(goals)), size=9.5)
            c = sum(1 for g in goals if g.confirm_status == "confirmed")
            r = sum(1 for g in goals if g.confirm_status == "rejected")
            pn = sum(1 for g in goals if g.confirm_status == "pending")
            _run(cells[3].paragraphs[0], f"{c} / {r} / {pn}", size=9.5)
            if d["refl"]:
                _run(cells[4].paragraphs[0], _fmt_dt(d["refl"].completed_at), color=OK, bold=True, size=9.5)
            else:
                _run(cells[4].paragraphs[0], "ещё нет", color=GREY, size=9.5)
            parts = []
            for g in goals:
                gr = d["grefl"].get(g.id)
                key = "irrelevant" if g.status == "dropped" else (gr.outcome if gr else None)
                parts.append({"achieved": "✓", "partial": "±", "not_achieved": "✗", "irrelevant": "—"}.get(key, "·"))
            _run(cells[5].paragraphs[0], " ".join(parts) or "—", size=9.5)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        _para(doc, "Обозначения в колонке «Итоги»: ✓ получилось · ± частично · ✗ не получилось · "
                   "— потеряла актуальность · · рефлексии ещё нет.", italic=True, color=GREY, size=9)

        no_goals = [d["name"] for d in datas if not d["goals"]]
        no_refl = [d["name"] for d in datas if d["goals"] and not d["refl"]]
        if no_goals:
            _para(doc, "Ничего не прислали: " + ", ".join(no_goals) + ".", color=BAD, size=10)
        if no_refl:
            _para(doc, "Поставили цели, но ещё не подвели итоги: " + ", ".join(no_refl) + ".", color=MID, size=10)

    # --- карты студентов ---
    doc.add_heading("2. Карты целей по студентам", level=1)
    for n, d in enumerate(datas, 1):
        if n > 1:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(f"2.{n}. {d['name']}", level=1)
        await _render_student(doc, session, d, mentor=mentor_user, goals_deadline=goals_deadline,
                              reflection_deadline=reflection_deadline, title_level=2, prefix=f"2.{n}.")
    _footer(doc)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue(), f"TutorAI_общий_отчёт_{_safe_name(mentor_name)}.docx", len(datas)
